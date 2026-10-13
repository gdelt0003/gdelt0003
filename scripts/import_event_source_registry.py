"""Import the editable Source Directory sheet from the event review workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook


def _enabled(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().casefold() in {"true", "yes", "1", "enabled"}


def import_registry(workbook_path: Path, seed_path: Path) -> dict[str, Any]:
    seed = json.loads(seed_path.read_text(encoding="utf-8"))
    if not isinstance(seed, dict) or not isinstance(seed.get("sources"), list):
        raise ValueError("Seed source registry must contain a sources list.")
    if not workbook_path.is_file():
        return seed

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Source Directory" not in workbook.sheetnames:
        workbook.close()
        return seed
    sheet = workbook["Source Directory"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        workbook.close()
        return seed
    columns = {str(value).strip(): index for index, value in enumerate(headers) if value}
    required = {"Category", "Source", "URL"}
    if not required.issubset(columns):
        workbook.close()
        raise ValueError(
            "Source Directory needs Category, Source, and URL columns."
        )

    seed_by_url = {
        source["url"]: source for source in seed["sources"] if source.get("url")
    }
    imported = []
    seen_ids: set[str] = set()
    seen_urls: set[str] = set()
    for row in rows:
        values = list(row)

        def get(name: str, default: Any = "") -> Any:
            index = columns.get(name)
            return values[index] if index is not None and index < len(values) else default

        url = str(get("URL") or "").strip()
        name = str(get("Source") or "").strip()
        category = str(get("Category") or "").strip()
        if not url and not name and not category:
            continue
        parsed = urlparse(url)
        if parsed.scheme != "https" or not parsed.netloc:
            raise ValueError(f"Source row must contain a valid HTTPS URL: {url!r}")
        if not name or not category:
            raise ValueError("Every non-empty source row requires Source and Category.")
        if url in seen_urls:
            raise ValueError(f"Duplicate source URL in workbook: {url}")

        original = seed_by_url.get(url, {})
        source_id = str(get("Source ID") or original.get("source_id") or "").strip()
        if not source_id:
            slug = re.sub(r"[^a-z0-9]+", "-", name.casefold()).strip("-")[:36]
            suffix = hashlib.sha256(url.encode("utf-8")).hexdigest()[:8]
            source_id = f"{slug or 'source'}-{suffix}"
        if source_id in seen_ids:
            raise ValueError(f"Duplicate source ID in workbook: {source_id}")
        source = {
            "source_id": source_id,
            "category": category,
            "name": name,
            "url": url,
            "enabled": _enabled(get("Enabled", original.get("enabled", True))),
            "collection_mode": str(
                get("Collection mode", original.get("collection_mode", "manual_review"))
                or "manual_review"
            ).strip(),
            "frequency": str(
                get("Target frequency", original.get("frequency", "daily")) or "daily"
            ).strip(),
            "last_attempt_utc": get(
                "Last attempt UTC", original.get("last_attempt_utc")
            ) or None,
            "last_result": str(
                get("Last result", original.get("last_result", "not_connected"))
                or "not_connected"
            ).strip(),
            "access_status": str(
                get(
                    "Access status",
                    original.get("access_status", "not_independently_checked"),
                )
                or "not_independently_checked"
            ).strip(),
            "added_by": str(get("Added by", original.get("added_by", "user")) or "user").strip(),
            "notes": str(get("Notes", original.get("notes", "")) or "").strip(),
            "last_run_url": str(
                get("Last run URL", original.get("last_run_url", "")) or ""
            ).strip(),
            "user_source_annotation_verbatim": str(
                get(
                    "User annotation on this source (verbatim)",
                    original.get("user_source_annotation_verbatim", ""),
                )
                or ""
            ).strip(),
        }
        if original or source.get("added_by") == "user":
            source["user_provided_url_verbatim"] = str(
                get("User-provided URL (verbatim)", url) or url
            ).strip()
        imported.append(source)
        seen_ids.add(source_id)
        seen_urls.add(url)

    if not imported:
        workbook.close()
        raise ValueError("Source Directory contains no source rows.")
    registry = {
        "schema_version": 1,
        "notes": seed.get("notes", ""),
        "user_source_list_instruction_verbatim": seed.get(
            "user_source_list_instruction_verbatim", ""
        ),
        "user_terminology_verbatim": seed.get("user_terminology_verbatim", ""),
        "user_terminology_clarification_verbatim": seed.get(
            "user_terminology_clarification_verbatim", ""
        ),
        "user_terminology_correction_verbatim": seed.get(
            "user_terminology_correction_verbatim", ""
        ),
        "user_field_expansion_permission_verbatim": seed.get(
            "user_field_expansion_permission_verbatim", ""
        ),
        "user_additional_source_request_verbatim": seed.get(
            "user_additional_source_request_verbatim", ""
        ),
        "sources": imported,
    }
    workbook.close()
    return registry


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--seed", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        registry = import_registry(args.workbook, args.seed)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(registry, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"Imported {len(registry['sources'])} event-discovery source rows.")
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
