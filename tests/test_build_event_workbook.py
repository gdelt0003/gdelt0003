import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from scripts.build_event_workbook import (
    build_workbook,
    load_catalog,
    load_source_registry,
)


class EventWorkbookTests(unittest.TestCase):
    def test_catalog_workbook_preserves_unknown_event_time_and_source_evidence(self):
        catalog_path = Path("outputs/event-catalog.json")
        catalog = load_catalog(catalog_path)
        source_registry = load_source_registry(
            Path("config/event-discovery-sources.json")
        )

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "event-sources.xlsx"
            build_workbook(catalog, output, source_registry=source_registry)
            workbook = load_workbook(output, read_only=False, data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Events",
                "Sources",
                "Source Directory",
                "Candidate Review",
                "Read me",
                "User Instructions",
                "Search Coverage",
            ],
        )
        events = list(workbook["Events"].values)
        headers = events[0]
        columbus = next(row for row in events[1:] if row[2] == "President Trump Hosts a Columbus Day Celebration")
        self.assertFalse(
            any(
                row[2]
                == "President Trump Participates in a Rose Garden Lunch, Oct. 15, 2026"
                for row in events[1:]
            )
        )
        self.assertIsNone(columbus[headers.index("Event date")])
        self.assertIsNone(columbus[headers.index("Event local time")])
        self.assertEqual(columbus[headers.index("Identity status")], "incomplete")
        self.assertGreaterEqual(columbus[headers.index("Source count")], 10)
        self.assertEqual(
            columbus[headers.index("President attendance duration (minutes; verified)")],
            None,
        )
        self.assertEqual(
            columbus[headers.index("Speech word count (verified; language/method)")],
            None,
        )
        instruction_rows = list(workbook["User Instructions"].values)
        instructions = " ".join(str(row[1]) for row in instruction_rows[1:])
        self.assertIn("为御注", instructions)
        self.assertIn("你们乐意添加", instructions)

        source_rows = list(workbook["Sources"].values)
        source_headers = source_rows[0]
        fox_news = next(row for row in source_rows[1:] if row[source_headers.index("Publisher")] == "Fox News")
        self.assertEqual(fox_news[source_headers.index("Video resolution")], "unknown")
        self.assertEqual(fox_news[source_headers.index("Frame review")], "sampled_at_20:01,25:01,27:01")
        self.assertEqual(fox_news[source_headers.index("Clean-source status")], "unverified")
        ap = next(
            row for row in source_rows[1:]
            if row[source_headers.index("Publisher")] == "Associated Press"
        )
        self.assertEqual(ap[source_headers.index("Video resolution")], "1920x1080 (HD1080 player decode)")
        self.assertEqual(ap[source_headers.index("Clean-source status")], "failed")
        self.assertIn("not watermark-free", ap[source_headers.index("Evidence and limitations")])

        candidate_rows = list(workbook["Candidate Review"].values)
        candidate_headers = candidate_rows[0]
        self.assertEqual(len(candidate_headers), len(set(candidate_headers)))
        self.assertTrue(
            all(len(row) == len(candidate_headers) for row in candidate_rows)
        )
        columbus_candidates = [
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Event title")]
            == "President Trump Hosts a Columbus Day Celebration"
        ]
        self.assertGreater(len(columbus_candidates), 0)
        white_house_anchor = next(
            row
            for row in columbus_candidates
            if row[candidate_headers.index("Publisher")] == "The White House"
        )
        self.assertIn(
            "embedded graphics reported, not a clean editing source",
            white_house_anchor[candidate_headers.index("Current source path")],
        )
        self.assertEqual(
            white_house_anchor[
                candidate_headers.index(
                    "Observed stream resolution (not thumbnail/source-master resolution)"
                )
            ],
            "unknown",
        )
        self.assertEqual(
            white_house_anchor[
                candidate_headers.index("Difference from official runtime (seconds)")
            ],
            0,
        )
        self.assertIn(
            "not quality rank",
            candidate_headers[candidate_headers.index("Category display slot (evidence order; not quality rank)")],
        )
        user_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=rQlpuY2tj-4"
        )
        self.assertIsNone(user_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            user_lead[candidate_headers.index("Category display slot (evidence order; not quality rank)")],
            "pending",
        )
        self.assertIn(
            "Do not mix with the October 13 Katie event",
            user_lead[candidate_headers.index("Current source path")],
        )
        self.assertEqual(
            user_lead[candidate_headers.index("Same-event evidence")],
            "pending_manual_review",
        )
        self.assertEqual(user_lead[candidate_headers.index("Duration (seconds)")], None)
        self.assertEqual(
            user_lead[
                candidate_headers.index(
                    "Observed stream resolution (not thumbnail/source-master resolution)"
                )
            ],
            "unknown",
        )
        self.assertEqual(
            user_lead[candidate_headers.index("Embedded graphics / occlusion review")],
            "pending_direct_frame_review",
        )
        self.assertEqual(
            user_lead[candidate_headers.index("Reuse rights review")],
            "pending_rights_review",
        )
        self.assertEqual(
            user_lead[
                candidate_headers.index("User-provided observation (internal note; unverified)")
            ],
            "这个源可以 更长 干净 高清 可信源,全体通报学习",
        )
        self.assertEqual(
            user_lead[
                candidate_headers.index("User subjective rating (not independent score)")
            ],
            90,
        )
        self.assertEqual(
            user_lead[
                candidate_headers.index("User angle opinion (not verified)")
            ],
            "打90分...此角度视频典范,还可以增补其它角度",
        )
        self.assertEqual(
            user_lead[candidate_headers.index("Camera location")], "unknown"
        )
        self.assertEqual(
            user_lead[
                candidate_headers.index("Media-seat annotation")
            ],
            "pending_verified_position",
        )
        self.assertEqual(
            user_lead[
                candidate_headers.index("Camera verification status")
            ],
            "pending",
        )
        rose_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=IySF_8OArvw"
        )
        self.assertIsNone(rose_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            rose_lead[
                candidate_headers.index(
                    "Category display slot (evidence order; not quality rank)"
                )
            ],
            "pending",
        )
        self.assertEqual(
            rose_lead[candidate_headers.index("Same-event evidence")],
            "official_video_page_title_and_date_claim; exact event/content/time match unverified",
        )
        self.assertEqual(
            rose_lead[candidate_headers.index("Source provenance / lineage (reported)")],
            "pending_source_provenance_review",
        )
        self.assertEqual(
            rose_lead[candidate_headers.index("Reuse rights review")],
            "pending_rights_review",
        )
        rose_evidence = rose_lead[
            candidate_headers.index("Evidence and limitations")
        ]
        self.assertIn(
            "User processing/publication authorization = confirmed",
            rose_evidence,
        )
        self.assertIn(
            "do not establish source-specific reuse rights",
            rose_evidence,
        )
        self.assertIn(
            "source lineage is unverified",
            rose_evidence,
        )
        self.assertIn("identify The White House and the title/date claim", rose_evidence)
        self.assertIn("do not establish direct-frame/audio content match, exact event time, or timezone", rose_evidence)
        self.assertIn("production_bbox_input=false", rose_evidence)
        self.assertIn("Current status (2026-10-16)", rose_evidence)
        self.assertIn("PR #50 merged as head 690c5b0 and merge commit 4613166", rose_evidence)
        self.assertIn("all six frontend, Python, and fixture-probe CI checks passed", rose_evidence)
        self.assertIn("Current main is 775c90e8", rose_evidence)
        self.assertIn("Production run 38868029978 used main SHA c6f3931", rose_evidence)
        self.assertIn("English cue 15 at 180.030s", rose_evidence)
        self.assertIn("Audit, finalization, artifact upload, and Pages deployment were skipped", rose_evidence)
        self.assertIn("Private artifact-only run 38895767927 on current main 775c90e8", rose_evidence)
        self.assertIn("render was in progress and artifact upload pending at last read", rose_evidence)
        self.assertIn("does not require the event-time/timezone or source-lineage URLs", rose_evidence)
        self.assertIn("logged both as empty and still passed validation", rose_evidence)
        self.assertIn("general https://www.whitehouse.gov/copyright/ page, not verified source-specific reuse evidence", rose_evidence)
        self.assertIn("operator confirmation do not establish source-specific reuse rights", rose_evidence)
        self.assertIn("source timecode/provenance=not_verified", rose_evidence)
        self.assertIn("No new successful artifact or production release is verified; production remains incomplete", rose_evidence)
        second_user_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=UgIivPOerNc"
        )
        self.assertIsNone(second_user_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            second_user_lead[candidate_headers.index("Category display slot (evidence order; not quality rank)")],
            "pending",
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("User subjective rating (not independent score)")
            ],
            50,
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("Repeated/static-scene review")
            ],
            "pending_direct_frame_review",
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("Source credibility review")
            ],
            "pending_source_review",
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("Independent camera review")
            ],
            "pending_source_lineage_review",
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("Cross-lead angle/lineage comparison")
            ],
            "pending_exact_event_angle_and_lineage_confirmation",
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index("User-provided observation (internal note; unverified)")
            ],
            "此链接50 分 虽然长 但是台标信号凑数 不变风景凑 还有logo 还有横幅遮挡",
        )
        third_user_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=Uya88MLbHV0"
        )
        self.assertIsNone(third_user_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            third_user_lead[candidate_headers.index("Category display slot (evidence order; not quality rank)")],
            "pending",
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index("User subjective rating (not independent score)")
            ],
            50,
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index(
                    "User feature score (not independent; rubric unspecified)"
                )
            ],
            10,
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index("User-provided observation (internal note; unverified)")
            ],
            "此时长一版 和白宫官方类似 横幅过大 但是特色是由不同角度 整体50分 其中特色分10分",
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index("Same-event evidence")
            ],
            "pending_exact_event_and_time_match",
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index("Cross-lead angle/lineage comparison")
            ],
            "pending_exact_event_match_people_stage_background_audio_continuity_and_source_lineage_review",
        )
        fourth_user_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=h30qlVplAsE"
        )
        self.assertIsNone(fourth_user_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            fourth_user_lead[candidate_headers.index("Category display slot (evidence order; not quality rank)")],
            "pending",
        )
        self.assertEqual(
            fourth_user_lead[
                candidate_headers.index("User subjective rating (not independent score)")
            ],
            30,
        )
        self.assertEqual(
            fourth_user_lead[
                candidate_headers.index("User-proposed role (not verified)")
            ],
            "user_proposed_background_reference_and_backup_candidate_only; not approved for production",
        )
        self.assertEqual(
            fourth_user_lead[
                candidate_headers.index("User-provided observation (internal note; unverified)")
            ],
            "此链接比官方还短 还有logo 30分,但是也可以作为背景 参考 剋有看看有没有不同角度 不同花絮,作为备选片子",
        )
        self.assertEqual(
            fourth_user_lead[candidate_headers.index("Same-event evidence")],
            "pending_exact_event_and_time_match",
        )
        self.assertEqual(
            fourth_user_lead[
                candidate_headers.index("Repeated/static-scene review")
            ],
            "pending_direct_frame_review",
        )
        self.assertEqual(
            fourth_user_lead[
                candidate_headers.index(
                    "Observed stream resolution (not thumbnail/source-master resolution)"
                )
            ],
            "unknown",
        )
        self.assertEqual(
            fourth_user_lead[candidate_headers.index("Reuse rights review")],
            "pending_rights_review",
        )
        self.assertIn(
            "genuinely different camera angles",
            fourth_user_lead[
                candidate_headers.index("User-requested follow-up (pending)")
            ],
        )
        low_score_lead = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.youtube.com/watch?v=rDw_no4Rjk8"
        )
        self.assertIsNone(low_score_lead[candidate_headers.index("Event ID")])
        self.assertEqual(
            low_score_lead[
                candidate_headers.index("User discovery classification (not verified)")
            ],
            "under_20_low_score_discovery; user classification, not verified score",
        )
        self.assertEqual(
            low_score_lead[
                candidate_headers.index("User subjective rating (not independent score)")
            ],
            None,
        )
        self.assertEqual(
            low_score_lead[
                candidate_headers.index("User-proposed role (not verified)")
            ],
            "discovery for checking same-event new angles and b-roll only; not a credibility finding and not production-ready",
        )
        self.assertFalse(
            low_score_lead[candidate_headers.index("Production-ready status")]
        )
        self.assertEqual(
            low_score_lead[
                candidate_headers.index("User-provided observation (internal note; unverified)")
            ],
            "shorter than official; few subscribers; logo/junk marks",
        )
        self.assertEqual(
            low_score_lead[
                candidate_headers.index("Neutral public wording if later verified")
            ],
            "publisher branding/station bug/graphics",
        )
        for url, expected_publisher, expected_title in (
            (
                "https://apnews.com/video/14-minutes-3-news-outlets-trumps-new-press-secretary-makes-abrupt-debut-93ce596f6efd415ead20d54aefdcf14f",
                "AP News",
                "14 minutes: 3 news outlets: Trump's new press secretary makes abrupt debut",
            ),
            (
                "https://www.kcentv.com/video/news/nation-world/first-briefing-from-new-white-house-press-secretary-katie-zacharia/507-3abb494f-a706-4bbd-b111-2e5283dcb981",
                "KCEN",
                "First briefing from new White House press secretary Katie Zacharia",
            ),
        ):
            supplementary = next(
                row
                for row in candidate_rows[1:]
                if row[candidate_headers.index("Source URL")] == url
            )
            self.assertIsNone(supplementary[candidate_headers.index("Event ID")])
            self.assertEqual(
                supplementary[candidate_headers.index("Publisher")],
                expected_publisher,
            )
            self.assertEqual(
                supplementary[candidate_headers.index("Candidate title")],
                expected_title,
            )
            self.assertEqual(
                supplementary[candidate_headers.index("Duration (seconds)")],
                None,
            )
            self.assertEqual(
                supplementary[
                    candidate_headers.index(
                        "Observed stream resolution (not thumbnail/source-master resolution)"
                    )
                ],
                "unknown",
            )
            self.assertEqual(
                supplementary[candidate_headers.index("Production-ready status")],
                False,
            )
        self.assertEqual(
            third_user_lead[candidate_headers.index("Duration (seconds)")], None
        )
        self.assertEqual(
            third_user_lead[
                candidate_headers.index(
                    "Observed stream resolution (not thumbnail/source-master resolution)"
                )
            ],
            "unknown",
        )
        self.assertEqual(
            second_user_lead[candidate_headers.index("Duration (seconds)")], None
        )
        self.assertEqual(
            second_user_lead[
                candidate_headers.index(
                    "Observed stream resolution (not thumbnail/source-master resolution)"
                )
            ],
            "unknown",
        )
        news_leads = [
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Event title")]
            == "Katie Zacharia briefing (Google News search leads; event and topic match unverified)"
        ]
        self.assertEqual(len(news_leads), 10)
        self.assertTrue(
            all(row[candidate_headers.index("Candidate category")] == "news" for row in news_leads)
        )
        self.assertEqual(
            {row[candidate_headers.index("Category display slot (evidence order; not quality rank)")] for row in news_leads},
            set(range(1, 11)),
        )
        video_search_leads = [
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Event title")]
            == "Oct 13 Katie briefing (search leads; event/source match unverified)"
            and row[candidate_headers.index("Candidate category")] == "video"
        ]
        self.assertEqual(len(video_search_leads), 10)
        self.assertTrue(
            all(
                row[candidate_headers.index("Top-10 display slot (not a result cap)")]
                == "Slots 1-10"
                for row in news_leads
            )
        )
        social_leads = [
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Event title")]
            == "Katie briefing social search leads (event/source match unverified)"
        ]
        self.assertEqual(len(social_leads), 10)
        self.assertTrue(
            all(row[candidate_headers.index("Candidate category")] == "social" for row in social_leads)
        )
        cspan_event_page = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Publisher")] == "C-SPAN"
            and row[candidate_headers.index("Source URL")]
            == "https://www.c-span.org/event/white-house-event/white-house-press-secretary-katie-zacharia-holds-first-briefing/448203"
        )
        cspan_facebook = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Source URL")]
            == "https://www.facebook.com/CSPAN/videos/white-house-press-secretary-katie-zacharia-told-reporters-on-tuesday-during-a-pr/1509770514326571/"
        )
        self.assertIsNone(cspan_event_page[candidate_headers.index("Event ID")])
        self.assertIsNone(cspan_facebook[candidate_headers.index("Event ID")])
        self.assertIn(
            "not reviewed",
            cspan_event_page[candidate_headers.index("Evidence and limitations")],
        )

        unlinked_report = next(
            row
            for row in candidate_rows[1:]
            if row[candidate_headers.index("Candidate title")]
            == "User-described screenshot; source unresolved"
        )
        self.assertIsNone(unlinked_report[candidate_headers.index("Source URL")])
        self.assertIsNone(unlinked_report[candidate_headers.index("Event ID")])
        self.assertEqual(
            unlinked_report[candidate_headers.index("Time marker present/status")],
            "user-reported; candidate unresolved",
        )
        self.assertEqual(
            unlinked_report[candidate_headers.index("Exact on-screen marker text")],
            "9:44 AM ET",
        )
        self.assertIsNone(
            unlinked_report[candidate_headers.index("Verified event time")]
        )
        self.assertEqual(
            unlinked_report[candidate_headers.index("Player chrome text (not event time)")],
            "0:10 / 14:07",
        )
        self.assertEqual(
            unlinked_report[candidate_headers.index("Unlinked screenshot overall score")],
            60,
        )
        self.assertEqual(
            unlinked_report[candidate_headers.index("Unlinked screenshot feature score")],
            15,
        )
        coverage_rows = list(workbook["Search Coverage"].values)
        coverage_headers = coverage_rows[0]
        coverage = {
            row[coverage_headers.index("Attempt ID")]: row
            for row in coverage_rows[1:]
        }
        self.assertEqual(len(coverage), 3)
        news_attempt = coverage["google-news-katie-briefing-2026-10-15"]
        self.assertEqual(news_attempt[coverage_headers.index("Results retained")], 10)
        self.assertIn(
            "not exhaustive",
            news_attempt[coverage_headers.index("Coverage limitations")],
        )
        for row in source_rows[1:]:
            local_time = row[source_headers.index("Page-published local display (not event time)")]
            timestamp = row[source_headers.index("Published metadata (UTC; not event time)")]
            if local_time and timestamp:
                parsed_local = datetime.strptime(local_time, "%Y-%m-%d %H:%M PDT")
                expected_utc = parsed_local.replace(
                    tzinfo=ZoneInfo("America/Los_Angeles")
                ).astimezone(timezone.utc)
                self.assertEqual(
                    datetime.fromisoformat(timestamp),
                    expected_utc,
                    f"publication timestamp mismatch for {row[source_headers.index('Source URL')]}",
                )

    def test_source_directory_and_user_context_are_included(self):
        catalog = load_catalog(Path("outputs/event-catalog.json"))
        source_registry = load_source_registry(
            Path("config/event-discovery-sources.json")
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "event-sources.xlsx"
            build_workbook(catalog, output, source_registry=source_registry)
            workbook = load_workbook(output, read_only=False, data_only=True)

        directory_rows = list(workbook["Source Directory"].values)
        source_headers = directory_rows[0]
        white_house = next(
            row for row in directory_rows[1:]
            if row[source_headers.index("Source ID")] == "white-house-youtube"
        )
        self.assertEqual(
            white_house[source_headers.index("Collection mode")],
            "automated_metadata",
        )
        gdelt = next(
            row for row in directory_rows[1:]
            if row[source_headers.index("Source ID")] == "gdelt-trump-articles"
        )
        self.assertEqual(
            gdelt[source_headers.index("Collection mode")],
            "automated_metadata",
        )
        self.assertEqual(
            white_house[source_headers.index("User-provided URL (verbatim)")],
            "https://www.youtube.com/@WhiteHouse",
        )
        self.assertEqual(
            white_house[source_headers.index("User annotation on this source (verbatim)")],
            None,
        )
        self.assertEqual(len(directory_rows) - 1, 13)
        additional_source = next(
            row for row in directory_rows[1:]
            if row[source_headers.index("Source ID")] == "forth-trump-social-lists"
        )
        self.assertEqual(
            additional_source[source_headers.index("Collection mode")],
            "manual_review",
        )
        instructions = " ".join(
            str(row[1]) for row in workbook["User Instructions"].values
        )
        self.assertIn("https://www.forth.news/lists/trumpts", instructions)

    def test_gdelt_articles_are_exported_as_metadata_only_leads(self):
        catalog = load_catalog(Path("outputs/event-catalog.json"))
        articles = {
            "checked_at_utc": "2026-10-13T12:00:00+00:00",
            "leads": [
                {
                    "title": "Recent Trump article",
                    "publisher": "example.com",
                    "url": "https://example.com/article",
                    "seen_date_utc_metadata": "20261013T110000Z",
                    "language": "English",
                    "source_country": "United States",
                    "assignment_status": "unassigned_event_lead",
                    "evidence_status": "article_metadata_only",
                }
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "event-sources.xlsx"
            build_workbook(catalog, output, gdelt_discovery=articles)
            workbook = load_workbook(output, read_only=False, data_only=True)

        self.assertIn("GDELT Article Leads", workbook.sheetnames)
        rows = list(workbook["GDELT Article Leads"].values)
        headers = rows[0]
        lead = rows[1]
        self.assertIsNone(lead[headers.index("Event ID")])
        self.assertEqual(lead[headers.index("Evidence status")], "article_metadata_only")
        self.assertIn("not verify event occurrence", lead[headers.index("Source/search limitations")])
        workbook.close()

    def test_candidate_review_selects_only_a_verified_same_moment_alternate(self):
        catalog = {
            "events": [
                {
                    "event_id": "evt-example",
                    "identity": {"title": "Example event"},
                    "official_reference": {
                        "duration_seconds": 100,
                        "video_url": "https://www.youtube.com/watch?v=official",
                    },
                    "sources": [
                        {
                            "publisher": "The White House",
                            "source_role": "official_event_anchor",
                            "url": "https://www.youtube.com/watch?v=official",
                            "match_status": "verified_event_anchor",
                            "clean_source_status": "failed",
                            "video_resolution": "1280x720",
                            "duration_seconds": 100,
                            "evidence": "Embedded graphics visible.",
                        },
                        {
                            "publisher": "Example News",
                            "source_role": "same_event_source_lead",
                            "url": "https://www.youtube.com/watch?v=alternate",
                            "match_status": "same_speech_moment_observed",
                            "clean_source_status": "passed",
                            "video_resolution": "1920x1080",
                            "duration_seconds": 105,
                            "independence_status": "independent_publisher",
                            "evidence": "Same speech moment directly matched.",
                        },
                        {
                            "publisher": "Unverified Lead",
                            "source_role": "same_event_search_lead",
                            "url": "https://www.youtube.com/watch?v=lead",
                            "match_status": "unverified",
                            "clean_source_status": "unverified",
                            "video_resolution": "unknown",
                            "duration_seconds": 100,
                        },
                    ],
                }
            ],
            "matching_rule": "event evidence",
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "candidate-review.xlsx"
            build_workbook(catalog, output)
            workbook = load_workbook(output, read_only=False, data_only=True)

        rows = list(workbook["Candidate Review"].values)
        headers = rows[0]
        alternate = next(
            row
            for row in rows[1:]
            if row[headers.index("Publisher")] == "Example News"
        )
        anchor = next(
            row
            for row in rows[1:]
            if row[headers.index("Publisher")] == "The White House"
        )
        self.assertEqual(alternate[headers.index("Category display slot (evidence order; not quality rank)")], 2)
        self.assertIn(
            "Preferred alternate",
            alternate[headers.index("Current source path")],
        )
        self.assertIn(
            "Another candidate meets",
            anchor[headers.index("Current source path")],
        )
        workbook.close()

    def test_candidate_review_is_uncapped_and_top_ten_is_category_specific(self):
        sources = [
            {
                "publisher": "The White House",
                "source_role": "official_event_anchor",
                "source_category": "video",
                "url": "https://www.youtube.com/watch?v=official",
                "match_status": "verified_event_anchor",
                "clean_source_status": "failed",
                "video_resolution": "1920x1080",
                "duration_seconds": 100,
            }
        ]
        for category in ("video", "news"):
            for index in range(12):
                sources.append(
                    {
                        "publisher": f"{category} publisher {index}",
                        "source_role": "same_event_search_lead",
                        "source_category": category,
                        "url": f"https://example.com/{category}/{index}",
                        "match_status": "same_event_context_only",
                        "clean_source_status": "unverified",
                        "video_resolution": "unknown",
                    }
                )
        catalog = {
            "events": [
                {
                    "event_id": "evt-many",
                    "identity": {"title": "Many source leads"},
                    "official_reference": {"duration_seconds": 100},
                    "sources": sources,
                }
            ],
            "matching_rule": "event evidence",
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "many-candidates.xlsx"
            build_workbook(catalog, output)
            workbook = load_workbook(output, read_only=False, data_only=True)

        rows = list(workbook["Candidate Review"].values)
        headers = rows[0]
        for category in ("video", "news"):
            category_rows = [
                row
                for row in rows[1:]
                if row[headers.index("Candidate category")] == category
            ]
            expected_count = 13 if category == "video" else 12
            self.assertEqual(len(category_rows), expected_count)
            ranks = {
                row[headers.index("Category display slot (evidence order; not quality rank)")]
                for row in category_rows
            }
            self.assertEqual(ranks, set(range(1, expected_count + 1)))
            self.assertEqual(
                sum(
                    row[headers.index("Top-10 display slot (not a result cap)")]
                    == "Slots 1-10"
                    for row in category_rows
                ),
                10,
            )
            self.assertEqual(
                sum(
                    row[headers.index("Top-10 display slot (not a result cap)")]
                    == "Additional result"
                    for row in category_rows
                ),
                expected_count - 10,
            )
        workbook.close()


if __name__ == "__main__":
    unittest.main()
